"""Integrações externas: Excel (xlsx) e WhatsApp.

Excel  -> exportação de todos os registros da obra e importação de entregas de EPI.
WhatsApp -> envio via WhatsApp Cloud API (Meta). Sem credenciais, devolve um
            link wa.me (clique-para-conversar) com a mensagem pronta.
"""
from __future__ import annotations

import io
import os
import urllib.parse
from datetime import datetime, timedelta
from typing import Any

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------- EXCEL
SHEETS: list[tuple[str, str, list[tuple[str, str]]]] = [
    ("EPIs", "epis", [("worker_name", "Trabalhador"), ("worker_company", "Empresa"), ("worker_role", "Função"),
                      ("epi_type", "EPI"), ("delivery_date", "Entrega"), ("expiry_date", "Validade")]),
    ("Não Conformidades", "nonconformities", [("title", "Título"), ("description", "Descrição"), ("severity", "Severidade"),
                                              ("status", "Status"), ("location", "Local"), ("user_name", "Registrado por"),
                                              ("created_at", "Criado em"), ("resolved_at", "Resolvido em")]),
    ("RDOs", "rdos", [("date", "Data"), ("text", "Relato"), ("weather", "Clima"), ("workers_count", "Efetivo"),
                      ("user_name", "Responsável")]),
    ("Checklists", "checklists", [("service_type", "Serviço"), ("status", "Status"), ("user_name", "Responsável"),
                                  ("created_at", "Criado em")]),
    ("Qualidade FVS-FVM", "quality", [("type", "Tipo"), ("service", "Serviço"), ("location", "Local"),
                                      ("result", "Resultado"), ("notes", "Observações"), ("created_at", "Data")]),
    ("Inspeções", "inspections", [("location", "Local"), ("user_name", "Responsável"), ("created_at", "Data")]),
]

HEADER_FILL = PatternFill("solid", fgColor="FF5E00")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_workbook(rows_by_table: dict[str, list[dict]], obra_name: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, table, cols in SHEETS:
        ws = wb.create_sheet(title)
        ws.append([label for _, label in cols])
        for c in ws[1]:
            c.fill, c.font = HEADER_FILL, HEADER_FONT
        for row in rows_by_table.get(table, []):
            ws.append([row.get(k) for k, _ in cols])
        if table == "inspections":
            ws.cell(row=1, column=len(cols) + 1, value="Resumo IA").font = HEADER_FONT
            ws.cell(row=1, column=len(cols) + 1).fill = HEADER_FILL
            for i, row in enumerate(rows_by_table.get(table, []), start=2):
                ws.cell(row=i, column=len(cols) + 1, value=(row.get("analysis") or {}).get("summary", ""))
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 22
        ws.freeze_panes = "A2"
    info = wb.create_sheet("Sobre", 0)
    info.append(["Obra", obra_name])
    info.append(["Exportado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EPI_IMPORT_COLUMNS = {
    "trabalhador": "worker_name", "nome": "worker_name",
    "empresa": "worker_company",
    "função": "worker_role", "funcao": "worker_role",
    "epi": "epi_type",
    "entrega": "delivery_date", "data de entrega": "delivery_date",
    "validade (dias)": "expiry_days", "dias de validade": "expiry_days",
    "validade": "expiry_date",
}


def epi_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "EPIs"
    ws.append(["Trabalhador", "Empresa", "Função", "EPI", "Entrega", "Validade (dias)"])
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    ws.append(["João Pereira", "Construtora ABC", "Pedreiro", "Capacete", datetime.now().strftime("%Y-%m-%d"), 180])
    for i in range(1, 7):
        ws.column_dimensions[get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_date(v: Any) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def parse_epi_sheet(content: bytes) -> tuple[list[dict], list[str]]:
    """Retorna (registros_validos, erros)."""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["Planilha vazia"]
    header = [str(h or "").strip().lower() for h in rows[0]]
    mapping = {i: EPI_IMPORT_COLUMNS[h] for i, h in enumerate(header) if h in EPI_IMPORT_COLUMNS}
    if "worker_name" not in mapping.values() or "epi_type" not in mapping.values():
        return [], ["Colunas obrigatórias: 'Trabalhador' e 'EPI'. Baixe o modelo em Integrações."]
    records, errors = [], []
    for n, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        rec = {field: row[i] for i, field in mapping.items() if i < len(row)}
        if not rec.get("worker_name") or not rec.get("epi_type"):
            errors.append(f"Linha {n}: trabalhador e EPI são obrigatórios")
            continue
        delivery = _to_date(rec.get("delivery_date")) or datetime.now().strftime("%Y-%m-%d")
        expiry = _to_date(rec.get("expiry_date"))
        if not expiry:
            try:
                days = int(rec.get("expiry_days") or 180)
            except (TypeError, ValueError):
                days = 180
            expiry = (datetime.strptime(delivery, "%Y-%m-%d") + timedelta(days=days)).strftime("%Y-%m-%d")
        records.append({
            "worker_name": str(rec["worker_name"]).strip(),
            "worker_company": str(rec.get("worker_company") or "").strip(),
            "worker_role": str(rec.get("worker_role") or "").strip(),
            "epi_type": str(rec["epi_type"]).strip(),
            "delivery_date": delivery,
            "expiry_date": expiry,
        })
    return records, errors


# ------------------------------------------------------------------ WHATSAPP
def whatsapp_config() -> dict:
    return {
        "token": os.environ.get("WHATSAPP_TOKEN", "").strip(),
        "phone_number_id": os.environ.get("WHATSAPP_PHONE_NUMBER_ID", "").strip(),
        "default_to": os.environ.get("WHATSAPP_DEFAULT_TO", "").strip(),
        "api_version": os.environ.get("WHATSAPP_API_VERSION", "v21.0").strip(),
    }


def whatsapp_configured() -> bool:
    c = whatsapp_config()
    return bool(c["token"] and c["phone_number_id"])


def normalize_phone(phone: str) -> str:
    digits = "".join(ch for ch in (phone or "") if ch.isdigit())
    if digits and len(digits) <= 11:  # número brasileiro sem DDI
        digits = "55" + digits
    return digits


def wa_link(phone: str, message: str) -> str:
    to = normalize_phone(phone)
    base = f"https://wa.me/{to}" if to else "https://wa.me/"
    return f"{base}?text={urllib.parse.quote(message)}"


def send_whatsapp(phone: str, message: str) -> dict:
    """Envia pela Cloud API. Sem credenciais, retorna link wa.me."""
    cfg = whatsapp_config()
    to = normalize_phone(phone or cfg["default_to"])
    if not whatsapp_configured():
        return {"sent": False, "mode": "link", "link": wa_link(to, message),
                "detail": "WhatsApp Cloud API não configurada — abra o link para enviar manualmente."}
    if not to:
        return {"sent": False, "mode": "error", "detail": "Informe o número de destino (com DDD)."}
    url = f"https://graph.facebook.com/{cfg['api_version']}/{cfg['phone_number_id']}/messages"
    try:
        r = requests.post(
            url,
            headers={"Authorization": f"Bearer {cfg['token']}", "Content-Type": "application/json"},
            json={"messaging_product": "whatsapp", "to": to, "type": "text", "text": {"preview_url": False, "body": message[:4096]}},
            timeout=20,
        )
    except requests.RequestException as exc:
        return {"sent": False, "mode": "error", "detail": f"Falha de rede: {exc}", "link": wa_link(to, message)}
    if r.status_code >= 300:
        return {"sent": False, "mode": "error", "detail": f"API WhatsApp {r.status_code}: {r.text[:300]}", "link": wa_link(to, message)}
    return {"sent": True, "mode": "api", "response": r.json()}


# ------------------------------------------------- FUNCIONÁRIOS DA OBRA (EXCEL)
WORKER_IMPORT_COLUMNS = {
    "nome": "name", "funcionário": "name", "funcionario": "name", "trabalhador": "name", "colaborador": "name",
    "função": "role", "funcao": "role", "cargo": "role",
    "empresa": "company", "empreiteira": "company",
    "setor": "sector", "equipe": "sector",
    "admissão": "admission_date", "admissao": "admission_date", "data de admissão": "admission_date",
}


def workers_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Funcionários"
    ws.append(["Nome", "Função", "Empresa", "Setor", "Admissão"])
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    ws.append(["João Pereira", "Pedreiro", "Construtora ABC", "Alvenaria", datetime.now().strftime("%d/%m/%Y")])
    ws.append(["Maria Souza", "Eletricista", "Elétrica XYZ", "Instalações", datetime.now().strftime("%d/%m/%Y")])
    for i in range(1, 6):
        ws.column_dimensions[get_column_letter(i)].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_workers_sheet(content: bytes) -> tuple[list[dict], list[str]]:
    """Lê a planilha de funcionários. Retorna (registros, avisos)."""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["Planilha vazia"]
    header = [str(h or "").strip().lower() for h in rows[0]]
    mapping = {i: WORKER_IMPORT_COLUMNS[h] for i, h in enumerate(header) if h in WORKER_IMPORT_COLUMNS}
    if "name" not in mapping.values():
        return [], ["A planilha precisa de uma coluna 'Nome'. Baixe o modelo no cadastro da obra."]
    records, errors = [], []
    for n, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in row):
            continue
        rec = {field: row[i] for i, field in mapping.items() if i < len(row)}
        name = str(rec.get("name") or "").strip()
        if not name:
            errors.append(f"Linha {n}: nome vazio, ignorada")
            continue
        records.append({
            "name": name,
            "role": str(rec.get("role") or "").strip(),
            "company": str(rec.get("company") or "").strip(),
            "sector": str(rec.get("sector") or "").strip(),
            "admission_date": _to_date(rec.get("admission_date")),
        })
    return records, errors
