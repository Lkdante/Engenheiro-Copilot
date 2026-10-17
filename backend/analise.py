"""Análise de obra no cadastro.

Dois caminhos:
  • OBRA GERAL — planilhas e arquivos por área (pessoas, segurança/EPI/EPC, almoxarifado, planejamento,
    qualidade, documentação, custos/contratos) + projeto (AutoCAD DXF/DWG, PDF). O app lê os arquivos,
    extrai indicadores, sugere os dados da obra e gera um parecer por área.
  • PATOLOGIA DAS CONSTRUÇÕES — fotos do local analisadas por visão computacional (manifestações
    patológicas, gravidade GUT, causa provável, recomendação). Sem chave de IA as fotos ficam
    pendentes de validação humana — nada é inventado.

Os extratores são determinísticos (openpyxl, csv, pypdf, leitor DXF próprio). A IA (Claude) é usada,
quando configurada, apenas para o parecer textual e para a visão das fotos.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import datetime, date
from statistics import mean
from typing import Any, Optional

from openpyxl import load_workbook

# --------------------------------------------------------------------- ÁREAS
AREAS: dict[str, dict] = {
    "pessoas": {"label": "Pessoas e acesso", "hint": "Lista de funcionários, terceiros, controle de acesso"},
    "seguranca": {"label": "Segurança, EPIs e EPCs", "hint": "Entregas de EPI, validade, CA, EPCs, treinamentos"},
    "almoxarifado": {"label": "Almoxarifado", "hint": "Materiais, estoque, estoque mínimo, entradas e saídas"},
    "planejamento": {"label": "Planejamento e produção", "hint": "Cronograma, etapas, % executado, datas"},
    "qualidade": {"label": "Qualidade", "hint": "FVS, FVM, inspeções, não conformidades"},
    "documentacao": {"label": "Documentação", "hint": "ART, alvará, licenças, laudos, projetos aprovados"},
    "custos": {"label": "Custos e contratos", "hint": "Orçamento, medições, contratos, previsto x realizado"},
    "projeto": {"label": "Projeto (AutoCAD / PDF)", "hint": "Arquivos DXF/DWG, plantas e memoriais em PDF"},
}
PHOTO_CATEGORY = "foto"
ALL_CATEGORIES = set(AREAS) | {PHOTO_CATEGORY}

SPREADSHEET_EXT = {".xlsx", ".xlsm", ".csv"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".heic"}


def norm(s: Any) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def file_kind(filename: str) -> str:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext in SPREADSHEET_EXT:
        return "planilha"
    if ext == ".xls":
        return "xls"
    if ext == ".pdf":
        return "pdf"
    if ext == ".dxf":
        return "dxf"
    if ext == ".dwg":
        return "dwg"
    if ext in IMAGE_EXT:
        return "imagem"
    return "outro"


# ----------------------------------------------------------------- PLANILHAS
def _to_number(v: Any) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace("%", "").replace(" ", "")
    if not s:
        return None
    if "," in s:  # formato brasileiro 1.234,56
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _to_date(v: Any) -> Optional[str]:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19] if "%H" in fmt else s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def read_tables(content: bytes, filename: str) -> list[dict]:
    """Lê .xlsx/.xlsm/.csv → [{sheet, rows: [[...]]}] (máx. 5000 linhas por aba)."""
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)][:5000]
        return [{"sheet": "CSV", "rows": rows}]
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i >= 5000:
                break
            rows.append(list(r))
        out.append({"sheet": ws.title, "rows": rows})
    wb.close()
    return out


def _find_header(rows: list[list]) -> int:
    """Primeira linha (entre as 15 iniciais) com >= 2 textos não vazios."""
    for i, r in enumerate(rows[:15]):
        texts = [c for c in r if isinstance(c, str) and c.strip() and _to_number(c) is None]
        if len(texts) >= 2:
            return i
    return 0


# rótulos de "ficha" (rótulo numa célula, valor ao lado ou abaixo)
FIELD_LABELS: list[tuple[str, list[str]]] = [
    ("name", ["nome da obra", "obra", "empreendimento", "projeto"]),
    ("address", ["endereco", "localidade", "local da obra", "local", "localizacao"]),
    ("art", ["art", "n da art", "numero da art", "no da art", "art n"]),
    ("company", ["empresa responsavel", "construtora", "contratada", "empresa", "executora"]),
    ("start_date", ["data de inicio", "inicio da obra", "inicio", "data inicial"]),
    ("end_date", ["termino previsto", "previsao de termino", "data de termino", "termino", "data final", "fim previsto"]),
    ("progress", ["avanco fisico", "% executado", "estagio atual", "avanco", "progresso", "percentual executado", "% concluido"]),
    ("workers_count", ["n de funcionarios", "numero de funcionarios", "funcionarios", "efetivo", "efetivo total"]),
    ("size", ["porte", "porte da obra"]),
]


def _label_match(cell: str) -> Optional[str]:
    c = norm(cell).rstrip(":").strip()
    c = c.replace("º", "").replace("°", "").replace(".", "")
    for field, labels in FIELD_LABELS:
        if c in labels:
            return field
    return None


def _kv_fields(rows: list[list]) -> dict:
    found: dict[str, Any] = {}
    for i, r in enumerate(rows[:200]):
        filled = sum(1 for c in r if c not in (None, ""))
        if filled > 3:  # linha de cabeçalho de tabela, não é ficha "rótulo: valor"
            continue
        for j, cell in enumerate(r):
            if not isinstance(cell, str):
                continue
            field = _label_match(cell)
            if not field or field in found:
                continue
            cands = []
            if j + 1 < len(r):
                cands.append(r[j + 1])
            # valor abaixo do rótulo só quando o rótulo está sozinho na linha
            if filled == 1 and i + 1 < len(rows) and j < len(rows[i + 1]):
                cands.append(rows[i + 1][j])
            for v in cands:
                if v in (None, ""):
                    continue
                if field in ("start_date", "end_date"):
                    d = _to_date(v)
                    if d:
                        found[field] = d
                        break
                elif field in ("progress", "workers_count"):
                    n = _to_number(v)
                    if n is not None:
                        if field == "progress" and 0 < n <= 1 and not str(v).strip().endswith("%"):
                            n *= 100  # célula formatada como % no Excel
                        found[field] = round(n, 1) if field == "progress" else int(n)
                        break
                elif isinstance(v, str) and v.strip() and _label_match(v) is None:
                    found[field] = v.strip()
                    break
    return found


def _columns(header: list) -> dict[str, int]:
    return {norm(h): i for i, h in enumerate(header) if isinstance(h, str) and h.strip()}


def _col(cols: dict[str, int], *names: str) -> Optional[int]:
    for n in names:
        for k, i in cols.items():
            if k == n or k.startswith(n + " ") or k.startswith(n):
                return i
    return None


def _data_rows(rows: list[list], header_idx: int) -> list[list]:
    return [r for r in rows[header_idx + 1:] if any(c not in (None, "") for c in r)]


def analyze_spreadsheet(content: bytes, filename: str, category: str) -> dict:
    tables = read_tables(content, filename)
    today = date.today().isoformat()
    sheets, fields, metrics, highlights, warnings = [], {}, {}, [], []

    for t in tables:
        rows = t["rows"]
        if not rows:
            continue
        for k, v in _kv_fields(rows).items():
            fields.setdefault(k, v)
        h = _find_header(rows)
        header = rows[h] if rows else []
        cols = _columns(header)
        data = _data_rows(rows, h)
        sheets.append({"name": t["sheet"], "rows": len(data), "columns": [str(c) for c in header if c not in (None, "")][:20]})

        name_c = _col(cols, "nome", "funcionario", "colaborador", "trabalhador")
        # ---------------- PESSOAS
        if category == "pessoas" and name_c is not None:
            names = [r[name_c] for r in data if name_c < len(r) and r[name_c]]
            metrics["funcionarios"] = metrics.get("funcionarios", 0) + len(names)
            comp_c = _col(cols, "empresa", "empreiteira")
            if comp_c is not None:
                comps = {str(r[comp_c]).strip() for r in data if comp_c < len(r) and r[comp_c]}
                metrics["empresas"] = len(comps)
        # ---------------- SEGURANÇA / EPI
        if category == "seguranca":
            epi_c = _col(cols, "epi", "equipamento", "item")
            val_c = _col(cols, "validade", "vencimento")
            ca_c = _col(cols, "ca", "c.a", "certificado")
            if epi_c is not None:
                metrics["registros_epi"] = metrics.get("registros_epi", 0) + len(data)
            if val_c is not None:
                vencidos = sum(1 for r in data if val_c < len(r) and (_to_date(r[val_c]) or "9999") < today)
                metrics["epis_vencidos"] = metrics.get("epis_vencidos", 0) + vencidos
                if vencidos:
                    warnings.append(f"{vencidos} EPI(s) com validade vencida na aba '{t['sheet']}'")
            if epi_c is not None and ca_c is None:
                warnings.append(f"Aba '{t['sheet']}' sem coluna de CA (Certificado de Aprovação)")
        # ---------------- ALMOXARIFADO
        if category == "almoxarifado":
            item_c = _col(cols, "material", "item", "descricao", "produto")
            qtd_c = _col(cols, "quantidade", "qtd", "estoque atual", "saldo", "estoque")
            min_c = _col(cols, "estoque minimo", "minimo", "est minimo")
            if item_c is not None:
                metrics["itens_estoque"] = metrics.get("itens_estoque", 0) + len(data)
            if qtd_c is not None and min_c is not None and qtd_c != min_c:
                abaixo = [str(r[item_c]) if item_c is not None and item_c < len(r) else "?" for r in data
                          if max(qtd_c, min_c) < len(r) and (_to_number(r[qtd_c]) or 0) < (_to_number(r[min_c]) or 0)]
                metrics["abaixo_minimo"] = metrics.get("abaixo_minimo", 0) + len(abaixo)
                if abaixo:
                    warnings.append(f"{len(abaixo)} material(is) abaixo do estoque mínimo: {', '.join(abaixo[:5])}")
        # ---------------- PLANEJAMENTO
        if category == "planejamento":
            pct_c = _col(cols, "% executado", "% concluido", "executado", "avanco", "% realizado", "realizado", "progresso", "%")
            ini_c = _col(cols, "inicio", "data inicio", "data de inicio")
            fim_c = _col(cols, "termino", "fim", "data fim", "data de termino", "previsao")
            etapa_c = _col(cols, "etapa", "atividade", "tarefa", "servico", "fase")
            if etapa_c is not None:
                metrics["etapas"] = metrics.get("etapas", 0) + len(data)
            if pct_c is not None:
                vals = []
                for r in data:
                    if pct_c < len(r):
                        n = _to_number(r[pct_c])
                        if n is not None:
                            vals.append(n * 100 if 0 < n <= 1 and not str(r[pct_c]).strip().endswith("%") else n)
                if vals:
                    metrics["avanco_medio"] = round(mean(vals), 1)
                    fields.setdefault("progress", metrics["avanco_medio"])
                    atrasadas = 0
                    if fim_c is not None:
                        atrasadas = sum(1 for r in data if max(fim_c, pct_c) < len(r)
                                        and (_to_date(r[fim_c]) or "9999") < today and (_to_number(r[pct_c]) or 0) < 100)
                        if atrasadas:
                            warnings.append(f"{atrasadas} etapa(s) com término vencido e não concluídas")
                    metrics["etapas_atrasadas"] = atrasadas
            ds = [d for c in (ini_c,) if c is not None for d in (_to_date(r[c]) for r in data if c < len(r)) if d]
            de = [d for c in (fim_c,) if c is not None for d in (_to_date(r[c]) for r in data if c < len(r)) if d]
            if ds:
                fields.setdefault("start_date", min(ds))
            if de:
                fields.setdefault("end_date", max(de))
        # ---------------- QUALIDADE
        if category == "qualidade":
            res_c = _col(cols, "resultado", "status", "situacao", "conforme")
            if res_c is not None:
                vals = [norm(r[res_c]) for r in data if res_c < len(r) and r[res_c]]
                ok = sum(1 for v in vals if v.startswith(("aprov", "conforme", "ok", "sim")))
                nok = sum(1 for v in vals if v.startswith(("reprov", "nao conforme", "nc", "nok", "nao")))
                metrics["inspecoes"] = metrics.get("inspecoes", 0) + len(vals)
                metrics["aprovadas"] = metrics.get("aprovadas", 0) + ok
                metrics["reprovadas"] = metrics.get("reprovadas", 0) + nok
                if nok:
                    warnings.append(f"{nok} inspeção(ões) reprovada(s)/não conforme(s)")
        # ---------------- DOCUMENTAÇÃO
        if category == "documentacao":
            doc_c = _col(cols, "documento", "descricao", "tipo", "nome")
            val_c = _col(cols, "validade", "vencimento")
            if doc_c is not None:
                docs = [norm(r[doc_c]) for r in data if doc_c < len(r) and r[doc_c]]
                metrics["documentos"] = metrics.get("documentos", 0) + len(docs)
                for essencial in ("art", "alvara", "pcmso", "pgr"):
                    if any(essencial in d for d in docs):
                        highlights.append(f"{essencial.upper()} presente")
                    else:
                        metrics.setdefault("faltando", []).append(essencial.upper())
            if val_c is not None:
                venc = sum(1 for r in data if val_c < len(r) and (_to_date(r[val_c]) or "9999") < today)
                if venc:
                    warnings.append(f"{venc} documento(s) vencido(s)")
        # ---------------- CUSTOS
        if category == "custos":
            tot_c = _col(cols, "valor total", "total", "valor", "custo total", "custo", "preco total", "orcado", "previsto")
            real_c = _col(cols, "realizado", "medido", "executado", "pago")
            if tot_c is not None:
                tot = sum(_to_number(r[tot_c]) or 0 for r in data if tot_c < len(r)
                          and not norm(r[0] if r else "").startswith("total"))
                metrics["orcamento_total"] = round(metrics.get("orcamento_total", 0) + tot, 2)
            if real_c is not None and real_c != tot_c:
                metrics["realizado_total"] = round(metrics.get("realizado_total", 0) +
                                                   sum(_to_number(r[real_c]) or 0 for r in data if real_c < len(r)), 2)
            # procura "TOTAL GERAL" escrito na planilha
            for r in rows:
                for j, c in enumerate(r):
                    if isinstance(c, str) and norm(c).startswith("total geral"):
                        nums = [_to_number(x) for x in r[j + 1:] if _to_number(x) is not None]
                        if nums:
                            metrics["orcamento_total"] = round(nums[-1], 2)

    # ART no formato texto em qualquer célula
    if "art" not in fields:
        for t in tables:
            for r in t["rows"][:300]:
                for c in r:
                    m = re.search(r"\bART\s*(?:n[ºo°.]*)?\s*[:\-]?\s*([A-Z]{0,2}[-\s]?\d[\d\.\-/]{5,})", str(c or ""), re.I)
                    if m:
                        fields["art"] = m.group(1).strip()
                        break
    return {"type": "planilha", "sheets": sheets, "fields": fields, "metrics": metrics,
            "highlights": highlights, "warnings": warnings}


# ----------------------------------------------------------------------- PDF
def analyze_pdf(content: bytes) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError:
        return {"type": "pdf", "warnings": ["Leitor de PDF não instalado (pip install pypdf)"]}
    try:
        reader = PdfReader(io.BytesIO(content))
        texts = []
        for page in reader.pages[:30]:
            try:
                texts.append(page.extract_text() or "")
            except Exception:
                texts.append("")
        text = "\n".join(texts)
    except Exception as exc:
        return {"type": "pdf", "warnings": [f"PDF ilegível: {exc}"]}
    rows = [[ln.split(":", 1)[0], ln.split(":", 1)[1]] if ":" in ln else [ln] for ln in text.splitlines()]
    fields = _kv_fields(rows)
    m = re.search(r"\bART\s*(?:n[ºo°.]*)?\s*[:\-]?\s*([A-Z]{0,2}[-\s]?\d[\d\.\-/]{5,})", text, re.I)
    if m and "art" not in fields:
        fields["art"] = m.group(1).strip()
    areas = re.findall(r"(?:[áa]rea\s+(?:constru[íi]da|total)[^\d]{0,20})([\d\.,]+)\s*m", text, re.I)
    metrics = {"paginas": len(reader.pages), "caracteres": len(text)}
    if areas:
        metrics["area_m2"] = _to_number(areas[0])
    warnings = [] if text.strip() else ["PDF sem texto (provavelmente digitalizado) — envie a versão editável ou fotos"]
    return {"type": "pdf", "fields": fields, "metrics": metrics, "snippet": text[:600], "warnings": warnings, "highlights": []}


# ----------------------------------------------------------------------- DXF
def analyze_dxf(content: bytes) -> dict:
    """Leitor simples de DXF (ASCII): camadas, contagem de entidades e textos do desenho."""
    text = content.decode("utf-8", errors="replace")
    if "SECTION" not in text[:5000]:
        return {"type": "dxf", "warnings": ["DXF binário ou inválido — salve como DXF ASCII no AutoCAD"]}
    lines = text.splitlines()
    pairs = [(lines[i].strip(), lines[i + 1].rstrip("\r")) for i in range(0, len(lines) - 1, 2)]
    entities: dict[str, int] = {}
    layers: set[str] = set()
    texts: list[str] = []
    in_entities = False
    current = None
    for code, value in pairs:
        if code == "2" and value.strip() == "ENTITIES":
            in_entities = True
            continue
        if code == "0" and value.strip() == "ENDSEC":
            in_entities = False
        if not in_entities:
            if code == "2" and current == "LAYER":
                layers.add(value.strip())
            if code == "0":
                current = value.strip()
            continue
        if code == "0":
            current = value.strip()
            entities[current] = entities.get(current, 0) + 1
        elif code == "8":
            layers.add(value.strip())
        elif code in ("1", "3") and current in ("TEXT", "MTEXT", "ATTRIB") and value.strip():
            clean = re.sub(r"\\[A-Za-z][^;]*;|[{}]|\\P", " ", value).strip()
            if clean and len(texts) < 300:
                texts.append(clean)
    joined = " | ".join(texts)
    pav = sorted({t for t in texts if re.search(r"pavimento|t[ée]rreo|subsolo|cobertura", t, re.I)})[:20]
    areas = [a for a in re.findall(r"([\d\.,]+)\s*m[²2]", joined)][:10]
    fields = {}
    m = re.search(r"\bART\s*[:\-]?\s*([\w\-/\.]{6,})", joined, re.I)
    if m:
        fields["art"] = m.group(1)
    return {
        "type": "dxf",
        "metrics": {"entidades": sum(entities.values()), "camadas": len(layers), "textos": len(texts)},
        "entities": dict(sorted(entities.items(), key=lambda kv: -kv[1])[:12]),
        "layers": sorted(layers)[:40],
        "pavimentos": pav,
        "areas_m2": areas,
        "texts_sample": texts[:25],
        "fields": fields,
        "highlights": [f"{len(layers)} camadas e {sum(entities.values())} entidades no desenho"] + ([f"Pavimentos: {', '.join(pav[:5])}"] if pav else []),
        "warnings": [],
    }


def analyze_file(content: bytes, filename: str, category: str) -> dict:
    kind = file_kind(filename)
    try:
        if kind == "planilha":
            return analyze_spreadsheet(content, filename, category)
        if kind == "pdf":
            return analyze_pdf(content)
        if kind == "dxf":
            return analyze_dxf(content)
    except Exception as exc:
        return {"type": kind, "warnings": [f"Não foi possível ler o arquivo: {exc}"], "fields": {}, "metrics": {}}
    if kind == "dwg":
        return {"type": "dwg", "fields": {}, "metrics": {}, "highlights": ["Arquivo DWG armazenado na obra"],
                "warnings": ["DWG é um formato fechado: para o app ler o desenho, exporte também em DXF (AutoCAD → Salvar como → DXF)"]}
    if kind == "xls":
        return {"type": "xls", "fields": {}, "metrics": {}, "warnings": ["Formato .xls antigo: salve como .xlsx no Excel para o app ler"]}
    if kind == "imagem":
        return {"type": "imagem", "fields": {}, "metrics": {}, "highlights": ["Imagem anexada"], "warnings": []}
    return {"type": "outro", "fields": {}, "metrics": {}, "warnings": ["Formato não suportado para leitura — arquivo apenas armazenado"]}


# ----------------------------------------------------------- PATOLOGIA (FOTO)
PATOLOGIA_SYSTEM = (
    "Você é engenheiro civil especialista em patologia das construções (NBR 15575, NBR 9575, NBR 6118, IBAPE). "
    "Analise a foto e retorne APENAS JSON válido no formato: "
    '{"manifestacoes":[{"tipo":"fissura|trinca|rachadura|umidade/infiltracao|eflorescencia|bolor/mofo|corrosao de armadura|'
    'desplacamento|destacamento de revestimento|segregacao/ninhos|recalque|deformacao excessiva|outro",'
    '"localizacao":"...","elemento":"laje|viga|pilar|alvenaria|fachada|piso|cobertura|outro","descricao":"...",'
    '"causa_provavel":"...","gravidade":1,"urgencia":1,"tendencia":1,"recomendacao":"...","confianca":"alta|media|baixa",'
    '"box_2d":[ymin,xmin,ymax,xmax]}],'
    '"estado_geral":"bom|regular|ruim|critico","ambiente":"...","summary":"...","requires_human_validation":true}. '
    "Gravidade, urgência e tendência vão de 1 a 5 (matriz GUT). box_2d é a caixa da manifestação na imagem, com "
    "coordenadas inteiras normalizadas de 0 a 1000 na ordem [ymin, xmin, ymax, xmax]. "
    "Nunca afirme manifestação sem evidência visual; "
    "se a imagem for inconclusiva use confianca baixa e diga 'Imagem inconclusiva — fotografe mais de perto e com boa luz'."
)


def gut(m: dict) -> int:
    try:
        return int(m.get("gravidade", 1)) * int(m.get("urgencia", 1)) * int(m.get("tendencia", 1))
    except (TypeError, ValueError):
        return 1


def consolidate_patologia(photo_results: list[dict]) -> dict:
    analyzed = [p for p in photo_results if p.get("status") == "analisado"]
    pend = len(photo_results) - len(analyzed)
    manif = []
    for i, p in enumerate(analyzed):
        for m in p.get("manifestacoes", []) or []:
            manif.append({**m, "gut": gut(m), "foto": i + 1})
    manif.sort(key=lambda m: -m["gut"])
    tipos: dict[str, int] = {}
    for m in manif:
        tipos[m.get("tipo", "outro")] = tipos.get(m.get("tipo", "outro"), 0) + 1
    max_gut = manif[0]["gut"] if manif else 0
    if not analyzed:
        estado = "pendente"
    elif max_gut >= 64:
        estado = "critico"
    elif max_gut >= 27:
        estado = "ruim"
    elif max_gut >= 8:
        estado = "regular"
    else:
        estado = "bom"
    return {
        "fotos": len(photo_results), "fotos_analisadas": len(analyzed), "fotos_pendentes": pend,
        "manifestacoes": manif[:50], "por_tipo": tipos, "gut_maximo": max_gut, "estado_geral": estado,
        "prioridades": [{"tipo": m.get("tipo"), "local": m.get("localizacao"), "gut": m["gut"], "acao": m.get("recomendacao")} for m in manif[:5]],
    }


# --------------------------------------------------------- OBRA GERAL (TUDO)
def suggest_size(workers: Optional[int], budget: Optional[float]) -> Optional[str]:
    if budget:
        return "pequeno" if budget < 1_500_000 else "medio" if budget < 20_000_000 else "grande"
    if workers:
        return "pequeno" if workers < 20 else "medio" if workers < 100 else "grande"
    return None


def consolidate_geral(files: list[dict]) -> dict:
    """files: [{category, filename, extracted}] → visão por área + campos sugeridos."""
    areas = {}
    fields: dict[str, Any] = {}
    # ordem de confiança para os dados da obra: fichas/documentos primeiro, depois cronograma e custos
    priority = ["documentacao", "planejamento", "custos", "projeto", "pessoas", "seguranca", "almoxarifado", "qualidade"]
    for key in priority:
        for f in files:
            if f["category"] == key:
                for k, v in ((f.get("extracted") or {}).get("fields") or {}).items():
                    fields.setdefault(k, v)
    for key, meta in AREAS.items():
        fs = [f for f in files if f["category"] == key]
        metrics: dict[str, Any] = {}
        highlights, warnings = [], []
        readable = 0
        for f in fs:
            ex = f.get("extracted") or {}
            if ex.get("type") in ("planilha", "pdf", "dxf"):
                readable += 1
            for k, v in (ex.get("metrics") or {}).items():
                if isinstance(v, (int, float)) and isinstance(metrics.get(k), (int, float)):
                    metrics[k] = round(metrics[k] + v, 2)
                elif isinstance(v, list):
                    metrics[k] = sorted(set(metrics.get(k, [])) | set(v))
                else:
                    metrics.setdefault(k, v)
            highlights += ex.get("highlights") or []
            warnings += ex.get("warnings") or []
        if key == "documentacao" and metrics.get("faltando"):
            present = {h.split(" ")[0] for h in highlights}
            falt = [d for d in metrics["faltando"] if d not in present]
            metrics["faltando"] = falt
            if falt:
                warnings.append("Não encontrado na documentação: " + ", ".join(falt))
        completeness = 0 if not fs else (100 if readable and metrics else 60 if readable else 30)
        status = "sem_dados" if not fs else ("atencao" if warnings else "ok")
        areas[key] = {"label": meta["label"], "arquivos": len(fs), "completude": completeness, "status": status,
                      "metricas": metrics, "destaques": highlights[:8], "alertas": warnings[:8]}

    pess = areas["pessoas"]["metricas"].get("funcionarios")
    if pess:
        fields["workers_count"] = int(pess)
    budget = areas["custos"]["metricas"].get("orcamento_total")
    size = suggest_size(fields.get("workers_count"), budget)
    if size and "size" not in fields:
        fields["size"] = size
    if isinstance(fields.get("size"), str):
        s = norm(fields["size"])
        fields["size"] = "pequeno" if s.startswith("peq") else "medio" if s.startswith("med") else "grande" if s.startswith("gra") else None
        if not fields["size"]:
            fields.pop("size")
    covered = sum(1 for a in areas.values() if a["arquivos"])
    score = round(sum(a["completude"] for a in areas.values()) / len(areas))
    pendencias = [a["label"] for a in areas.values() if not a["arquivos"]]
    return {"areas": areas, "campos_sugeridos": fields, "areas_cobertas": covered, "indice_completude": score,
            "pendencias": pendencias,
            "alertas": [f"{a['label']}: {w}" for a in areas.values() for w in a["alertas"]][:20]}


def local_parecer_geral(c: dict) -> str:
    lines = [f"Índice de completude do cadastro: {c['indice_completude']}% ({c['areas_cobertas']} de {len(AREAS)} áreas com arquivos)."]
    for key, a in c["areas"].items():
        if not a["arquivos"]:
            continue
        m = a["metricas"]
        parts = []
        if key == "pessoas" and m.get("funcionarios"):
            parts.append(f"{m['funcionarios']} funcionários" + (f" de {m['empresas']} empresa(s)" if m.get("empresas") else ""))
        if key == "seguranca" and "registros_epi" in m:
            parts.append(f"{m['registros_epi']} registros de EPI, {m.get('epis_vencidos', 0)} vencidos")
        if key == "almoxarifado" and "itens_estoque" in m:
            parts.append(f"{m['itens_estoque']} itens, {m.get('abaixo_minimo', 0)} abaixo do mínimo")
        if key == "planejamento" and "avanco_medio" in m:
            parts.append(f"avanço médio {m['avanco_medio']}%, {m.get('etapas_atrasadas', 0)} etapa(s) atrasada(s)")
        if key == "qualidade" and "inspecoes" in m:
            parts.append(f"{m['inspecoes']} inspeções ({m.get('aprovadas', 0)} aprovadas, {m.get('reprovadas', 0)} reprovadas)")
        if key == "custos" and m.get("orcamento_total"):
            parts.append(f"orçamento R$ {m['orcamento_total']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        if key == "projeto" and m.get("entidades"):
            parts.append(f"desenho com {m['entidades']} entidades em {m.get('camadas', 0)} camadas")
        lines.append(f"• {a['label']}: " + ("; ".join(parts) if parts else f"{a['arquivos']} arquivo(s) anexado(s)"))
    if c["pendencias"]:
        lines.append("Áreas sem arquivos: " + ", ".join(c["pendencias"]) + ".")
    if c["alertas"]:
        lines.append("Pontos de atenção: " + "; ".join(c["alertas"][:6]) + ".")
    return "\n".join(lines)


def local_parecer_patologia(c: dict) -> str:
    if not c["fotos"]:
        return "Nenhuma foto enviada."
    if not c["fotos_analisadas"]:
        return (f"{c['fotos']} foto(s) registradas e pendentes de análise: a visão computacional não está configurada "
                "(ANTHROPIC_API_KEY). Faça a vistoria visual e registre as manifestações manualmente.")
    lines = [f"{c['fotos_analisadas']} de {c['fotos']} foto(s) analisadas. Estado geral estimado: {c['estado_geral'].upper()} "
             f"(maior GUT = {c['gut_maximo']})."]
    if c["por_tipo"]:
        lines.append("Manifestações: " + ", ".join(f"{k} ({v})" for k, v in c["por_tipo"].items()) + ".")
    for p in c["prioridades"][:3]:
        lines.append(f"• Prioridade GUT {p['gut']}: {p['tipo']} — {p.get('local') or 'local não informado'}. {p.get('acao') or ''}")
    lines.append("Resultado de apoio à decisão: confirme em vistoria técnica presencial antes de emitir laudo.")
    return "\n".join(lines)


# ------------------------------------------------------ MODELOS DE PLANILHA
TEMPLATES: dict[str, list[list]] = {
    "pessoas": [["Nome", "Função", "Empresa", "Setor", "Admissão"], ["João Pereira", "Pedreiro", "Construtora ABC", "Alvenaria", "01/03/2026"]],
    "seguranca": [["Funcionário", "EPI", "CA", "Entrega", "Validade"], ["João Pereira", "Capacete", "31469", "01/03/2026", "01/09/2026"]],
    "almoxarifado": [["Material", "Unidade", "Quantidade", "Estoque mínimo", "Fornecedor"], ["Cimento CP-II 50kg", "saco", 120, 50, "Fornecedor X"]],
    "planejamento": [["Etapa", "Início", "Término", "% Executado", "Responsável"], ["Fundação", "01/03/2026", "30/04/2026", 100, "Eng. Ana"]],
    "qualidade": [["Serviço", "Local", "Data", "Resultado", "Observação"], ["Alvenaria", "Pav. 2", "10/04/2026", "Aprovado", ""]],
    "documentacao": [["Nome da obra", "Residencial Exemplo"], ["Endereço", "Rua X, 100 - Cidade/UF"], ["Construtora", "Construtora ABC"],
                     ["ART", "SP-2026-000000"], ["Data de início", "01/03/2026"], ["Término previsto", "30/12/2027"], [],
                     ["Documento", "Número", "Validade"], ["ART", "SP-2026-000000", ""], ["Alvará de construção", "123/2026", "31/12/2027"],
                     ["PGR", "", "01/03/2027"], ["PCMSO", "", "01/03/2027"]],
    "custos": [["Item", "Descrição", "Valor total", "Realizado"], ["1", "Fundação", 300000, 300000], ["2", "Estrutura", 900000, 250000],
               ["", "TOTAL GERAL", 1200000, 550000]],
}


def template_xlsx(area: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = AREAS[area]["label"][:31]
    for r in TEMPLATES[area]:
        ws.append(r)
    head = 7 if area == "documentacao" else 0
    for c in ws[head + 1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="FF5E00")
    for col in "ABCDE":
        ws.column_dimensions[col].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
