"""Testes de ponta a ponta da API local (SQLite temporário, sem chaves de IA).

Rodar:  cd backend && python -m pytest -q
"""
import io
import os
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    os.environ["SQLITE_PATH"] = str(tmp_path_factory.mktemp("db") / "test.db")
    os.environ["JWT_SECRET"] = "test-secret"
    os.environ["ANTHROPIC_API_KEY"] = ""
    os.environ["OPENAI_API_KEY"] = ""
    os.environ["WHATSAPP_TOKEN"] = ""
    os.environ["WHATSAPP_PHONE_NUMBER_ID"] = ""
    os.environ["GEMINI_API_KEY"] = ""
    os.environ["GROQ_API_KEY"] = ""
    from fastapi.testclient import TestClient
    import server

    with TestClient(server.app) as c:
        yield c


def login(client, email="engenheiro@demo.com", password="demo123"):
    r = client.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["database"] == "sqlite"


def test_login_ok_and_me(client):
    h = login(client)
    me = client.get("/api/auth/me", headers=h).json()
    assert me["email"] == "engenheiro@demo.com"
    assert "password_hash" not in me
    assert me["obra_id"]


def test_login_wrong_password_rejected(client):
    r = client.post("/api/auth/login", json={"email": "engenheiro@demo.com", "password": "errada"})
    assert r.status_code == 401


def test_login_unknown_user_rejected(client):
    r = client.post("/api/auth/login", json={"email": "ninguem@x.com", "password": "qualquer"})
    assert r.status_code == 401


def test_protected_routes_need_token(client):
    assert client.get("/api/dashboard").status_code == 401
    assert client.get("/api/dashboard", headers={"Authorization": "Bearer lixo"}).status_code == 401


def test_register_then_login(client):
    import re
    import emailer
    r = client.post("/api/auth/register", json={"name": "Novo Usuário", "email": "novo@obra.com", "password": "segredo1", "role": "almoxarife"})
    assert r.status_code == 200, r.text
    assert r.json()["pending"] is True and "token" not in r.json()
    # antes de confirmar o e-mail a conta não existe
    assert client.post("/api/auth/login", json={"email": "novo@obra.com", "password": "segredo1"}).status_code == 401
    corpo = [m for m in emailer.SENT if m["to"] == "novo@obra.com"][-1]["body"]
    token = re.search(r"token=([\w-]+)", corpo).group(1)
    assert "Conta criada" in client.get(f"/api/auth/verify-email?token={token}").text
    assert client.post("/api/auth/register", json={"name": "Dup", "email": "novo@obra.com", "password": "segredo1"}).status_code == 400
    h = login(client, "novo@obra.com", "segredo1")
    assert client.get("/api/auth/me", headers=h).json()["obra_id"]


def test_dashboard_and_alerts(client):
    h = login(client)
    d = client.get("/api/dashboard", headers=h).json()
    assert d["obra"]["name"] == "Residencial Vila Nova"
    assert d["kpis"]["nc_open"] == 4
    alerts = client.get("/api/alerts", headers=h).json()["items"]
    assert any(a["type"] == "nc" for a in alerts)


def test_copilot_local_mode_uses_real_data(client):
    h = login(client)
    r = client.post("/api/copilot/chat", headers=h, json={"session_id": "s1", "message": "Existe alguma não conformidade crítica aberta?"})
    assert r.status_code == 200
    body = r.json()
    assert body["mode"] == "local"
    assert "guarda-corpo" in body["message"]
    hist = client.get("/api/copilot/history/s1", headers=h).json()["messages"]
    assert [m["role"] for m in hist] == ["user", "assistant"]


def test_checklist_flow(client):
    h = login(client)
    c = client.post("/api/checklists/generate", headers=h, json={"service_type": "eletrica"}).json()
    for it in c["items"]:
        c = client.post(f"/api/checklists/{c['id']}/update-item", headers=h, json={"item_id": it["id"], "status": "ok"}).json()
    assert c["status"] == "concluido"


def test_rdo_epi_quality_nc(client):
    h = login(client)
    assert client.post("/api/rdo", headers=h, json={"text": "Concretagem da laje"}).status_code == 200
    e = client.post("/api/epis", headers=h, json={"worker_name": "Zé", "epi_type": "Capacete", "delivery_date": "2026-01-01", "expiry_days": 30}).json()
    assert e["alert"] is True
    assert client.post("/api/quality", headers=h, json={"type": "FVS", "service": "Pintura", "result": "aprovado"}).status_code == 200
    nc = client.post("/api/nc", headers=h, json={"title": "Teste", "severity": "baixa"}).json()
    assert client.post(f"/api/nc/{nc['id']}/resolve", headers=h).json()["ok"]


def test_photo_without_ai_is_saved_pending(client):
    h = login(client)
    p = client.post("/api/photos", headers=h, json={"image_base64": "aGVsbG8="}).json()
    assert p["analysis"]["status"] == "pendente"
    assert p["analysis"]["requires_human_validation"] is True


def test_transcribe_without_key_returns_503(client):
    h = login(client)
    r = client.post("/api/rdo/transcribe", headers=h, files={"file": ("a.m4a", b"123", "audio/m4a")})
    assert r.status_code == 503


def test_excel_export_and_import(client):
    from openpyxl import load_workbook
    h = login(client)
    r = client.get("/api/integrations/excel/export", headers=h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "EPIs" in wb.sheetnames and wb["EPIs"].max_row > 1
    tpl = client.get("/api/integrations/excel/epi-template", headers=h).content
    r = client.post("/api/integrations/excel/import-epis", headers=h, files={"file": ("m.xlsx", tpl, "application/octet-stream")})
    assert r.json()["imported"] == 1


def test_whatsapp_fallback_link(client):
    h = login(client)
    r = client.post("/api/integrations/whatsapp/alerts", headers=h, json={"to": "11999998888"}).json()
    assert r["sent"] is False and r["link"].startswith("https://wa.me/5511999998888")


# ----------------------------------------------------------------- OBRAS
NOVA_OBRA = {
    "name": "Edifício Aurora", "address": "Rua das Flores, 100 - Campinas/SP", "start_date": "01/03/2026",
    "end_date": "30/11/2027", "progress": 15, "company": "Construtora Aurora", "art": "SP-2026-777",
    "size": "grande", "workers_count": 30,
}


def test_obra_create_and_list(client):
    h = login(client)
    r = client.post("/api/obras", headers=h, json=NOVA_OBRA)
    assert r.status_code == 200, r.text
    o = r.json()
    assert o["start_date"] == "2026-03-01" and o["end_date"] == "2027-11-30"
    assert o["workers_total"] == 30
    items = client.get("/api/obras", headers=h).json()["items"]
    assert any(i["id"] == o["id"] for i in items)


def test_obra_validation(client):
    h = login(client)
    bad = {**NOVA_OBRA, "end_date": "01/01/2026"}
    assert client.post("/api/obras", headers=h, json=bad).status_code == 422
    assert client.post("/api/obras", headers=h, json={**NOVA_OBRA, "progress": 150}).status_code == 422
    assert client.post("/api/obras", headers=h, json={**NOVA_OBRA, "size": "gigante"}).status_code == 422


def test_obra_create_forbidden_for_estagiario(client):
    h = login(client, "estagiario@demo.com")
    assert client.post("/api/obras", headers=h, json=NOVA_OBRA).status_code == 403


def test_active_obra_header_isolates_data(client):
    h = login(client)
    o = client.post("/api/obras", headers=h, json={**NOVA_OBRA, "name": "Obra Isolada"}).json()
    ho = {**h, "X-Obra-Id": o["id"]}
    d = client.get("/api/dashboard", headers=ho).json()
    assert d["obra"]["name"] == "Obra Isolada"
    assert d["kpis"]["nc_open"] == 0
    client.post("/api/nc", headers=ho, json={"title": "NC da obra nova", "severity": "alta"})
    assert len(client.get("/api/nc", headers=ho).json()["items"]) == 1
    assert client.get("/api/dashboard", headers={**h, "X-Obra-Id": "nao-existe"}).status_code == 404


def test_obra_update(client):
    h = login(client)
    o = client.post("/api/obras", headers=h, json=NOVA_OBRA).json()
    u = client.put(f"/api/obras/{o['id']}", headers=h, json={"progress": 40, "end_date": "15/12/2027"}).json()
    assert u["progress"] == 40 and u["end_date"] == "2027-12-15"


def test_workers_import_from_excel(client):
    h = login(client)
    o = client.post("/api/obras", headers=h, json={**NOVA_OBRA, "workers_count": 0}).json()
    tpl = client.get("/api/integrations/excel/workers-template", headers=h).content
    r = client.post(f"/api/obras/{o['id']}/workers/import", headers=h, files={"file": ("f.xlsx", tpl, "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 2
    assert client.get(f"/api/obras/{o['id']}", headers=h).json()["workers_total"] == 2
    ws = client.get(f"/api/obras/{o['id']}/workers", headers=h).json()["items"]
    assert {w["name"] for w in ws} == {"João Pereira", "Maria Souza"}


# ------------------------------------------------------------ MENU PÚBLICO
def test_public_obras_without_login(client):
    r = client.get("/api/public/obras")
    assert r.status_code == 200
    items = r.json()["items"]
    vila = next(i for i in items if i["name"] == "Residencial Vila Nova")
    assert vila["art"] and vila["address"] and vila["company"]
    # dados internos não vazam no menu público
    assert "nc_open" not in vila and "created_by" not in vila and "workers_registered" not in vila


def test_public_obras_search(client):
    r = client.get("/api/public/obras", params={"q": "vila"}).json()["items"]
    assert r and all("vila" in i["name"].lower() for i in r)
    assert client.get("/api/public/obras", params={"q": "zzz-nao-existe"}).json()["items"] == []


def test_public_obra_detail_and_private_routes_still_protected(client):
    oid = client.get("/api/public/obras").json()["items"][0]["id"]
    assert client.get(f"/api/public/obras/{oid}").status_code == 200
    assert client.get("/api/dashboard", headers={"X-Obra-Id": oid}).status_code == 401
    assert client.get(f"/api/obras/{oid}").status_code == 401


# ----------------------------------------------- CADASTRO COM ANÁLISE DE ARQUIVOS
def _xlsx(rows_by_sheet: dict) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in rows_by_sheet.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DXF_SAMPLE = "\n".join([
    "0", "SECTION", "2", "ENTITIES",
    "0", "LINE", "8", "PAREDES", "10", "0", "20", "0", "11", "10", "21", "0",
    "0", "LINE", "8", "PAREDES", "10", "10", "20", "0", "11", "10", "21", "10",
    "0", "TEXT", "8", "TEXTOS", "1", "PAVIMENTO TIPO",
    "0", "TEXT", "8", "TEXTOS", "1", "AREA CONSTRUIDA 1.250,00 m2",
    "0", "ENDSEC", "0", "EOF", ""])


def _upload(client, h, oid, category, name, content):
    r = client.post(f"/api/obras/{oid}/arquivos", headers=h, params={"category": category},
                    files={"file": (name, content, "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def test_cadastro_geral_por_planilhas(client):
    h = login(client)
    d = client.post("/api/obras/rascunho", headers=h, json={"kind": "geral"}).json()
    oid = d["id"]
    assert d["status"] == "rascunho" and d["kind"] == "geral"
    # rascunho não aparece no menu público
    assert all(i["id"] != oid for i in client.get("/api/public/obras").json()["items"])

    ficha = _xlsx({"Ficha": [["Nome da obra", "Torre Horizonte"], ["Endereço", "Av. Brasil, 500 - Curitiba/PR"],
                             ["Construtora", "Horizonte Engenharia"], ["ART", "PR-2026-445566"],
                             ["Data de início", "01/02/2026"], ["Término previsto", "30/06/2027"]]})
    f1 = _upload(client, h, oid, "documentacao", "ficha.xlsx", ficha)
    assert f1["extracted"]["fields"]["art"] == "PR-2026-445566"

    pessoas = _xlsx({"Funcionários": [["Nome", "Função", "Empresa"], ["Ana", "Pedreira", "A"], ["Bia", "Servente", "A"], ["Caio", "Eletricista", "B"]]})
    _upload(client, h, oid, "pessoas", "pessoas.xlsx", pessoas)
    plan = _xlsx({"Cronograma": [["Etapa", "Início", "Término", "% Executado"], ["Fundação", "01/02/2026", "30/04/2026", 100],
                                 ["Estrutura", "01/05/2026", "30/12/2026", 40]]})
    _upload(client, h, oid, "planejamento", "cronograma.xlsx", plan)
    custos = _xlsx({"Orçamento": [["Item", "Valor total"], ["Fundação", 300000], ["Estrutura", 900000], ["TOTAL GERAL", 1200000]]})
    _upload(client, h, oid, "custos", "orcamento.xlsx", custos)
    estoque = _xlsx({"Estoque": [["Material", "Quantidade", "Estoque mínimo"], ["Cimento", 10, 50], ["Areia", 30, 10]]})
    _upload(client, h, oid, "almoxarifado", "estoque.xlsx", estoque)
    dxf = _upload(client, h, oid, "projeto", "planta.dxf", DXF_SAMPLE.encode())
    assert dxf["extracted"]["metrics"]["entidades"] == 4 and "PAVIMENTO TIPO" in dxf["extracted"]["pavimentos"]
    dwg = _upload(client, h, oid, "projeto", "planta.dwg", b"AC1032binario")
    assert dwg["extracted"]["warnings"]

    r = client.post(f"/api/obras/{oid}/analisar", headers=h)
    assert r.status_code == 200, r.text
    a = r.json()
    res = a["analysis"]["resultado"]
    assert res["areas"]["pessoas"]["metricas"]["funcionarios"] == 3
    assert res["areas"]["planejamento"]["metricas"]["avanco_medio"] == 70
    assert res["areas"]["almoxarifado"]["metricas"]["abaixo_minimo"] == 1
    assert res["areas"]["custos"]["metricas"]["orcamento_total"] == 1200000
    assert a["analysis"]["parecer"]
    o = a["obra"]
    assert o["name"] == "Torre Horizonte" and o["company"] == "Horizonte Engenharia"
    assert o["start_date"] == "2026-02-01" and o["end_date"] == "2027-06-30"
    assert o["progress"] == 70 and o["workers_total"] == 3 and o["size"] == "pequeno"

    c = client.post(f"/api/obras/{oid}/confirmar", headers=h)
    assert c.status_code == 200, c.text
    pub = next(i for i in client.get("/api/public/obras").json()["items"] if i["id"] == oid)
    assert pub["kind"] == "geral" and "analysis" not in pub


def test_cadastro_patologia_por_fotos(client):
    h = login(client)
    oid = client.post("/api/obras/rascunho", headers=h, json={"kind": "patologia"}).json()["id"]
    assert client.post(f"/api/obras/{oid}/analisar", headers=h).status_code == 422  # sem fotos
    png = bytes.fromhex("89504E470D0A1A0A0000000D4948445200000001000000010806000000")
    foto = _upload(client, h, oid, "foto", "fissura.png", png)
    assert foto["extracted"]["status"] == "pendente"  # sem chave de IA: não inventa resultado
    r = client.post(f"/api/obras/{oid}/analisar", headers=h).json()
    assert r["analysis"]["resultado"]["estado_geral"] == "pendente"
    # confirmar exige os dados de classificação
    assert client.post(f"/api/obras/{oid}/confirmar", headers=h).status_code == 422
    client.put(f"/api/obras/{oid}", headers=h, json={**NOVA_OBRA, "name": "Laudo Ed. Primavera"})
    assert client.post(f"/api/obras/{oid}/confirmar", headers=h).status_code == 200
    pub = next(i for i in client.get("/api/public/obras").json()["items"] if i["id"] == oid)
    assert pub["kind"] == "patologia"


def test_rascunho_cancelar_e_validacoes(client):
    h = login(client)
    assert client.post("/api/obras/rascunho", headers=h, json={"kind": "xyz"}).status_code == 422
    assert client.post("/api/obras/rascunho", headers=login(client, "estagiario@demo.com"), json={"kind": "geral"}).status_code == 403
    oid = client.post("/api/obras/rascunho", headers=h, json={"kind": "geral"}).json()["id"]
    assert client.post(f"/api/obras/{oid}/arquivos", headers=h, params={"category": "foto"},
                       files={"file": ("a.xlsx", b"x", "application/octet-stream")}).status_code == 422
    assert client.delete(f"/api/obras/{oid}", headers=h).status_code == 200
    assert client.get(f"/api/obras/{oid}", headers=h).status_code == 404


def test_modelos_por_area_sao_lidos_pelo_analisador(client):
    import analise
    h = login(client)
    for area in analise.TEMPLATES:
        r = client.get(f"/api/analise/modelo/{area}", headers=h)
        assert r.status_code == 200
        ex = analise.analyze_spreadsheet(r.content, "m.xlsx", area)
        assert ex["sheets"], area
    doc = analise.analyze_spreadsheet(client.get("/api/analise/modelo/documentacao", headers=h).content, "m.xlsx", "documentacao")
    assert doc["fields"]["art"] == "SP-2026-000000" and doc["fields"]["name"] == "Residencial Exemplo"


# ------------------------------------------------------ EXCLUSÃO DE OBRA (ADMIN)
def test_admin_exclui_obra_com_todos_os_dados(client):
    h_eng = login(client)
    h_adm = login(client, "admin@demo.com")
    o = client.post("/api/obras", headers=h_eng, json={**NOVA_OBRA, "name": "Obra Para Excluir"}).json()
    ho = {**h_eng, "X-Obra-Id": o["id"]}
    client.post("/api/nc", headers=ho, json={"title": "NC", "severity": "alta"})
    client.post("/api/rdo", headers=ho, json={"text": "RDO"})
    client.post(f"/api/obras/{o['id']}/workers", headers=h_eng, json={"name": "Fulano"})

    # engenheiro não pode excluir obra cadastrada
    assert client.delete(f"/api/obras/{o['id']}", headers=h_eng, params={"confirm": "Obra Para Excluir"}).status_code == 403
    # admin precisa confirmar com o nome exato
    assert client.delete(f"/api/obras/{o['id']}", headers=h_adm, params={"confirm": "outro nome"}).status_code == 422
    r = client.delete(f"/api/obras/{o['id']}", headers=h_adm, params={"confirm": "Obra Para Excluir"})
    assert r.status_code == 200, r.text
    ap = r.json()["apagados"]
    assert ap["nonconformities"] == 1 and ap["rdos"] == 1 and ap["workers"] == 1
    assert client.get(f"/api/public/obras/{o['id']}").status_code == 404
    assert all(i["id"] != o["id"] for i in client.get("/api/public/obras").json()["items"])


# ------------------------------------------- DEMO DE PATOLOGIA + VALIDAÇÃO (YOLO) + VOZ
def test_obra_demo_laudo_ufsc_no_menu(client):
    items = client.get("/api/public/obras").json()["items"]
    demo = next(i for i in items if i.get("demo") and "UFSC" in i["name"])
    assert demo["kind"] == "patologia" and "UFSC" in demo["name"]
    assert demo["demo_fonte"]["url"].startswith("https://dfo.ufsc.br/")
    an = client.get(f"/api/public/obras/{demo['id']}").json()["demo_analysis"]
    assert an["resultado"]["estado_geral"] == "ruim" and len(an["resultado"]["manifestacoes"]) >= 10
    # obras normais não expõem a análise publicamente
    assert all("demo_analysis" not in i for i in items if not i.get("demo"))
    h = {**login(client), "X-Obra-Id": demo["id"]}
    assert client.get("/api/dashboard", headers=h).json()["kpis"]["nc_open"] >= 3


def test_validacao_humana_e_status_do_dataset(client):
    import db
    h = login(client)
    oid = client.post("/api/obras/rascunho", headers=h, json={"kind": "patologia"}).json()["id"]
    png = bytes.fromhex("89504E470D0A1A0A0000000D4948445200000001000000010806000000")
    f = _upload(client, h, oid, "foto", "f.png", png)
    # simula uma análise com caixa (como o Gemini devolveria)
    ex = {**f["extracted"], "status": "analisado", "manifestacoes": [{"tipo": "fissura", "gut": 27, "box_2d": [100, 100, 500, 500]}]}
    db.update("obra_files", f["id"], {"extracted": ex})
    r = client.put(f"/api/obras/{oid}/arquivos/{f['id']}/validacao", headers=h, json={"itens": [{"index": 0, "status": "confirmada"}]})
    assert r.status_code == 200, r.text
    assert r.json()["extracted"]["validacao"]["0"]["status"] == "confirmada"
    assert client.put(f"/api/obras/{oid}/arquivos/{f['id']}/validacao", headers=h, json={"itens": [{"index": 5, "status": "confirmada"}]}).status_code == 422
    st = client.get("/api/analise/dataset/status", headers=h).json()
    assert st["confirmadas"] >= 1 and st["yolo"]["available"] is False
    assert client.get("/api/analise/dataset/yolo", headers=h).status_code == 403  # só admin
    z = client.get("/api/analise/dataset/yolo", headers=login(client, "admin@demo.com"))
    assert z.status_code == 200 and z.content[:2] == b"PK"


def test_copiloto_voz_sem_chave_retorna_503(client):
    h = login(client)
    r = client.post("/api/copilot/transcribe", headers=h, files={"file": ("v.m4a", b"123", "audio/m4a")})
    assert r.status_code == 503 and "GROQ_API_KEY" in r.json()["detail"]


def test_obras_reais_de_demonstracao(client):
    items = client.get("/api/public/obras").json()["items"]
    reais = {i["name"]: i for i in items if i.get("demo") and "UFSC" not in i["name"]}
    assert len(reais) == 4
    kinds = sorted(i["kind"] for i in reais.values())
    assert kinds == ["geral", "geral", "geral", "patologia"]
    concluidas = [i for i in reais.values() if i["kind"] == "geral" and i["progress"] >= 100]
    andamento = [i for i in reais.values() if i["kind"] == "geral" and i["progress"] < 100]
    assert len(concluidas) == 2 and len(andamento) == 1
    for o in reais.values():  # todos os campos que o cadastro exige
        for campo in ("name", "address", "start_date", "end_date", "company", "art", "size"):
            assert o[campo], (o["name"], campo)
        assert o["workers_total"] > 0 and 1 <= o["banner"] <= 8
        assert o["demo_fonte"]["url"].startswith("https://")
    huop = next(o for o in reais.values() if o["kind"] == "patologia")
    an = client.get(f"/api/public/obras/{huop['id']}").json()["demo_analysis"]
    assert an["resultado"]["estado_geral"] == "critico" and an["resultado"]["gut_maximo"] == 100
